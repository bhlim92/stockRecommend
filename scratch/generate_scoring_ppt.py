import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_scoring_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    
    # Theme colors
    bg_color = RGBColor(11, 14, 23)      # Slate Dark Mode (#0B0E17)
    card_color = RGBColor(22, 28, 45)    # Slightly lighter card background (#161C2D)
    text_white = RGBColor(248, 249, 250) # Off-white (#F8F9FA)
    text_gray = RGBColor(173, 181, 189)  # Cool gray (#ADB5BD)
    accent_cyan = RGBColor(0, 210, 255)  # Cyan (#00D2FF)
    accent_green = RGBColor(0, 255, 102) # Green (#00FF66)
    accent_red = RGBColor(255, 0, 85)    # Red (#FF0055)
    accent_gold = RGBColor(255, 193, 7)  # Amber (#FFC107)

    # Slide Master Layout Helper
    def apply_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

    def add_title(slide, text, subtitle_text=None):
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = accent_cyan
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Malgun Gothic"
            p2.font.size = Pt(14)
            p2.font.color.rgb = text_gray
            p2.space_before = Pt(4)

    def draw_card(slide, left, top, width, height, title, content_lines, accent_color=None):
        # Draw base card shape
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = card_color
        if accent_color:
            shape.line.color.rgb = accent_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.color.rgb = RGBColor(40, 50, 75)
            shape.line.width = Pt(0.5)

        # Text box inside card
        tf_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tf_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = accent_cyan if not accent_color else accent_color
        p.space_after = Pt(8)

        for line in content_lines:
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.name = "Malgun Gothic"
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = text_white
            p_line.space_before = Pt(4)
            
    # -------------------------------------------------------------
    # Slide 1: Cover Slide
    # -------------------------------------------------------------
    slide_layout = prs.slide_layouts[6] # Blank
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)

    # Accent decorative box
    line_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(2.2), Inches(0.15), Inches(3.2))
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = accent_cyan
    line_shape.line.fill.background()

    # Title box
    title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(3.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "주식 퀀트 AI 스크리너 알고리즘 설계"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = text_white
    p.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "진입 점수 (Entry Score) 및 평가 점수 (Evaluation Score) 산출 공식"
    p2.font.name = "Malgun Gothic"
    p2.font.size = Pt(20)
    p2.font.color.rgb = accent_cyan
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "주식 발굴 및 포트폴리오 리밸런싱 시스템 설계서 부록"
    p3.font.name = "Malgun Gothic"
    p3.font.size = Pt(12)
    p3.font.color.rgb = text_gray

    # -------------------------------------------------------------
    # Slide 2: Overview
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "1. 퀀트 평가 모델 개요 (System Overview)", "핵심 평가 엔진인 진입 점수와 평가 점수의 아키텍처적 역할 정의")

    # Left Card: Entry Score
    draw_card(slide, Inches(0.75), Inches(2.0), Inches(5.6), Inches(4.5), 
              "진입 점수 (Entry Score) - 기술적 시점 분석", [
                  "• 목적: 현재 주가가 기술적으로 매수 또는 진입하기에 적합한 시점인지 판정",
                  "• 분석 방법: 단기 추세, 모멘텀, 조정 깊이, 수급 조건을 계량 평가",
                  "• 핵심 지표: 이동평균선(SMA), MACD, RSI, 거래량(Volume)",
                  "• 가중치 합산 방식: 최종 실수 스펙(-1.0 ~ +1.0)을 산출한 뒤, 시스템 저장 및 UI 렌더링을 위해 0~100점 범위로 매핑하여 저장",
                  "• 권장 매매 행동 연동: 점수 범위에 따라 매수(Buy) / 보유(Hold) / 축소(Reduce) / 매도(Sell) 판단 자동 결정"
              ], accent_cyan)

    # Right Card: Evaluation Score
    draw_card(slide, Inches(6.98), Inches(2.0), Inches(5.6), Inches(4.5), 
              "평가 점수 (Evaluation Score) - 기본적 기업 가치 분석", [
                  "• 목적: 기업의 본질적인 재무 성장 모멘텀 및 저평가 상태 검증",
                  "• 분석 방법: 밸류에이션, 이익 성장률, 상승 여력, 명가 투자 규칙 점수화",
                  "• 핵심 지표: trailingPE, pegRatio, trailingEps, forwardEps, targetMeanPrice, CANSLIM 통과여부",
                  "• 가중치 배분: PER(30%) + PEG(20%) + EPS성장(20%) + 상승여력(15%) + CANSLIM(15%) = 총 100점 만점 구조",
                  "• 용도: 자산 배분 비중 조정 및 추천 리포트 최종 종목 선별 기준"
              ], accent_gold)

    # -------------------------------------------------------------
    # Slide 3: Entry Score Formulas
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "2. 진입 점수 (Entry Score) 산출 공식", "기술적 분석 4대 핵심 지표 신호 가중 합산 및 0~100점 선형 매핑")

    # Formula Box
    formula_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.8), Inches(11.83), Inches(1.2))
    formula_shape.fill.solid()
    formula_shape.fill.fore_color.rgb = card_color
    formula_shape.line.color.rgb = accent_cyan
    formula_shape.line.width = Pt(1.5)
    
    tf = formula_shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "최종 점수 (Final Score) = 0.4 × S1 (추세) + 0.3 × S2 (모멘텀) + 0.2 × S3 (눌림목) + 0.1 × S4 (수급)\n시스템 매핑 점수 (Mapped Score) = int((Final Score + 1.0) × 50)  [ 0점 ~ 100점 범위 ]"
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = text_white

    # 4 Signals Cards (Columns)
    col_w = Inches(2.78)
    gap = Inches(0.23)
    
    draw_card(slide, Inches(0.75), Inches(3.3), col_w, Inches(3.4), "S1. 추세 (이평선 정배열)\n[가중치 40%]", [
        "• 5일/20일/200일선 분석",
        "• +1점: 현재가 > 5 > 20 > 200 (정배열 상승)",
        "• -1점: 현재가 < 5 < 20 < 200 (역배열 하락)",
        "• 0점: 이평선 혼조세 및 횡보 구간"
    ])
    
    draw_card(slide, Inches(0.75 + 3.01), Inches(3.3), col_w, Inches(3.4), "S2. 모멘텀 (MACD 교차)\n[가중치 30%]", [
        "• MACD & 시그널선 크로스",
        "• +1점: MACD 0선 이상에서 시그널 상향 돌파 (골든크로스)",
        "• -1점: 시그널선 하향 돌파 (데드크로스)",
        "• 0점: 기타 일반 유지 구간"
    ])
    
    draw_card(slide, Inches(0.75 + 6.02), Inches(3.3), col_w, Inches(3.4), "S3. 눌림목 (RSI 조정)\n[가중치 20%]", [
        "• 과매수/과매도 분석",
        "• +1점: 상승장(S1=1)에서 RSI 40~50 조정 시 (건전한 진입 기회)",
        "• -1점: RSI 80 초과 과매수 구간 (단기 과열)",
        "• 0점: 기타 중간 안전 구간"
    ])

    draw_card(slide, Inches(0.75 + 9.03), Inches(3.3), col_w, Inches(3.4), "S4. 수급 (거래량 돌파)\n[가중치 10%]", [
        "• 거래 대금 수급 여력",
        "• +1점: 당일 거래량 > 최근 20일 평균의 1.5배 이상 & 당일 양봉 마감",
        "• 0점: 거래량이 미수반된 흐름 또는 음봉 마감"
    ])

    # -------------------------------------------------------------
    # Slide 4: Entry Score Decisions
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "3. 진입 점수에 따른 종합 매매 판단 (Action Map)", "산출된 Mapped Score(0~100)를 기준으로 권장 매매 판단 매칭")

    # Table creation
    rows, cols = 5, 3
    left, top, width, height = Inches(0.75), Inches(2.2), Inches(11.83), Inches(4.2)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # Set column widths
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(4.83)

    headers = ["최종 실수 점수 (Final Score)", "매핑 정수 점수 (Mapped Score)", "권장 매매 행동 (Action Strategy)"]
    for j, val in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = card_color
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Malgun Gothic"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = accent_cyan

    data = [
        ["+0.6 ~ +1.0", "80점 ~ 100점", "🟢 적극 매수 (Strong Buy) - 강한 상승 추세 및 안전 진입점 형성"],
        ["0.0 ~ +0.5", "50점 ~ 75점", "🟡 분할 매수 또는 보유 (Hold/Buy) - 추세 유지 중이나 추격 매수 유의"],
        ["-0.5 ~ -0.1", "25점 ~ 45점", "🟠 비중 축소 또는 관망 (Reduce/Wait) - 단기 하방 압력 또는 혼조"],
        ["-1.0 ~ -0.6", "0점 ~ 20점", "🔴 적극 매도 (Strong Sell) - 완전한 역배열 하락 추세 지속"]
    ]

    for i, row_data in enumerate(data):
        for j, val in enumerate(row_data):
            cell = table.cell(i+1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j < 2 else PP_ALIGN.LEFT
            p.font.name = "Malgun Gothic"
            p.font.size = Pt(13)
            p.font.color.rgb = text_white
            if "🟢" in val:
                p.font.color.rgb = accent_green
            elif "🔴" in val:
                p.font.color.rgb = accent_red
            elif "🟡" in val:
                p.font.color.rgb = accent_gold

    # -------------------------------------------------------------
    # Slide 5: Evaluation Score Overview
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "4. 평가 점수 (Evaluation Score) 배분", "기본적 재무 가치 및 성장성, 컨센서스, CANSLIM 만족도를 종합 합산(총 100점)")

    # Left: Weight Breakdown Table
    rows, cols = 6, 3
    left, top, width, height = Inches(0.75), Inches(2.2), Inches(6.5), Inches(4.2)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(1.2)
    table.columns[2].width = Inches(2.5)

    headers = ["평가 항목", "가중치(만점)", "yfinance 매핑 지표"]
    for j, val in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = card_color
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Malgun Gothic"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = accent_cyan

    eval_data = [
        ["① PER 밸류에이션 평가", "30점", "trailingPE"],
        ["② PEG 성장대비 주가 평가", "20점", "pegRatio"],
        ["③ 미래 EPS 성장성", "20점", "trailing/forwardEps"],
        ["④ 애널리스트 상승여력", "15점", "targetMeanPrice"],
        ["⑤ CANSLIM 필터 가산점", "15점", "passed_screener"],
    ]

    for i, row_data in enumerate(eval_data):
        for j, val in enumerate(row_data):
            cell = table.cell(i+1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j == 1 else PP_ALIGN.LEFT
            p.font.name = "Malgun Gothic"
            p.font.size = Pt(11)
            p.font.color.rgb = text_white
            if j == 1:
                p.font.bold = True
                p.font.color.rgb = accent_cyan

    # Right: Summary Card
    draw_card(slide, Inches(7.5), Inches(2.2), Inches(5.08), Inches(4.2), "기본적 분석 기반 평가 철학", [
        "• 재무 건전성 및 이익 성장 속도 결합:",
        "  전형적인 가치주 지표(PER)와 모멘텀 성장주 지표",
        "  (PEG, EPS성장, CANSLIM)를 적절히 혼합 설계",
        "",
        "• 데이터 수집의 한계 예방:",
        "  yfinance 등에서 특정 재무 정보 누락 시(예: 적자),",
        "  해당 항목만 0점 처리하고 전체 합산이 유지되도록",
        "  코드 레벨에서 유연한 예외 처리 설계 적용",
        "",
        "• 최종 종목 추천 랭킹 산출에 직결:",
        "  스크리너 화면에서 점수 기반 정렬에 최우선 사용"
    ], accent_gold)

    # -------------------------------------------------------------
    # Slide 6: Evaluation Score Formulas
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "5. 평가 항목별 상세 산출 로직", "각 재무 지표에 대한 구체적인 선형 감점 및 분기 기준 공식 명세")

    # Draw 4 cards for formulas
    cw, ch = Inches(5.6), Inches(2.0)
    
    # 1. PER
    draw_card(slide, Inches(0.75), Inches(2.2), cw, ch, "① PER (Trailing PE) 점수 - 만점 30점", [
        "• PE ≤ 10: 30점 만점 부여",
        "• 10 < PE < 30: int(30 × (30 - PE) / 20) 점 (선형 비례 감점)",
        "• PE ≥ 30 또는 적자(PE < 0)인 경우: 0점"
    ])
    
    # 2. PEG
    draw_card(slide, Inches(6.98), Inches(2.2), cw, ch, "② PEG (PEG Ratio) 점수 - 만점 20점", [
        "• PEG ≤ 1.0: 20점 만점 부여 (피터 린치 성장주 가치 기준)",
        "• 1.0 < PEG < 2.0: int(20 × (2.0 - PEG) / 1.0) 점 (선형 감점)",
        "• PEG ≥ 2.0 또는 데이터 누락 시: 0점"
    ])

    # 3. EPS Growth
    draw_card(slide, Inches(0.75), Inches(4.5), cw, ch, "③ EPS 성장성 점수 - 만점 20점", [
        "• 성장률 G = ((Forward EPS - Trailing EPS) / |Trailing EPS|) × 100",
        "• G ≥ 20%: 20점 만점 (미래 이익 고성장 모멘텀)",
        "• 0% < G < 20%: int(20 × G / 20) 점 (성장율에 비례한 점수)",
        "• G ≤ 0%: 0점"
    ])

    # 4. Target Price Space
    draw_card(slide, Inches(6.98), Inches(4.5), cw, ch, "④ 목표가 대비 상승 여력 - 만점 15점", [
        "• 상승여력 U = ((Target Mean Price - Current) / Current) × 100",
        "• U ≥ 20%: 15점 만점 (애널리스트 합의 20% 이상 저평가)",
        "• 0% < U < 20%: int(15 × U / 20) 점 (상승 여력 비례 점수)",
        "• U ≤ 0% 또는 정보 없음: 0점"
    ])

    # -------------------------------------------------------------
    # Slide 7: CANSLIM Rules
    # -------------------------------------------------------------
    slide = prs.slides.add_slide(slide_layout)
    apply_background(slide)
    add_title(slide, "6. CANSLIM 필터 세부 통과 기준 (15점 가산)", "윌리엄 오닐의 7대 모멘텀 투자 요건 (전부 만족 시 15점, 하나라도 미충족 시 0점)")

    # 7 Rules Cards
    card_w2 = Inches(3.7)
    card_h2 = Inches(2.1)
    
    # Row 1
    draw_card(slide, Inches(0.75), Inches(2.2), card_w2, card_h2, "C: 분기 EPS 성장 (Current)", [
        "• 최근 분기 주당순이익(EPS)이",
        "  전년 동기 대비(YoY) 20% 이상",
        "  상승했을 것"
    ])
    
    draw_card(slide, Inches(0.75 + 3.9), Inches(2.2), card_w2, card_h2, "A: 연간 이익 증가 (Annual)", [
        "• 최근 3개년 연간 EPS가 매년",
        "  20% 이상 성장하고, 최신 분기",
        "  ROE가 17% 이상일 것"
    ])

    draw_card(slide, Inches(0.75 + 7.8), Inches(2.2), card_w2, card_h2, "N: 신제품/신고가 (New)", [
        "• 기업의 혁신 동력이 있으며,",
        "  현재 주가가 52주 신고가",
        "  범위 인근(15% 이내)일 것"
    ])

    # Row 2
    draw_card(slide, Inches(0.75), Inches(4.6), card_w2, card_h2, "S: 수급과 거래량 (Supply)", [
        "• 당일 주가가 2% 이상 상승 시,",
        "  거래량이 50일 평균 거래량을",
        "  초과하여 강력 매수세 확인"
    ])

    draw_card(slide, Inches(0.75 + 3.9), Inches(4.6), card_w2, card_h2, "L/I: 주도주 & 기관 수급", [
        "• L: RS 백분위 순위 70 이상",
        "• I: 기관 투자자 지분율이",
        "  30% 이상일 것 (yfinance 기준)"
    ])

    draw_card(slide, Inches(0.75 + 7.8), Inches(4.6), card_w2, card_h2, "M: 시장 방향성 (Market)", [
        "• 소속 시장 종합 지수(KOSPI,",
        "  S&P500)가 50일선 위에 있고,",
        "  50일선이 200일선 위의 상승세"
    ])

    prs.save("design/screener_scoring_rules.pptx")
    print("PowerPoint saved to design/screener_scoring_rules.pptx successfully.")

if __name__ == '__main__':
    create_scoring_ppt()
